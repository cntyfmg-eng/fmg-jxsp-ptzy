# -*- coding: utf-8 -*-
"""生成《优质课配套教案》DOCX（含视频分析报告 + AI视频提示词）。

对应 Skill 步骤：调用「教师教案与课堂活动」后产出教案。
用法：
    python gen_lesson_plan.py [lesson_content.json]
不传参时默认读取同目录下的 lesson_content.json。
输出：./output/<课题>_优质课配套教案.docx
"""
import sys
import os
import json

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from docx_common import (new_doc, add_title, add_subtitle, add_h1, add_h2,
                         add_para, add_bullet, set_cell, style_table, header_row, shade,
                         NAVY, BLUE, GREY)

DEFAULT = os.path.join(HERE, "lesson_content.json")


def load(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def build(c, out_path):
    doc = new_doc()
    m = c["meta"]
    v = c.get("video", {})
    add_title(doc, "《%s》%s · 优质课配套教案" % (m["title"], m["grade"]))
    add_subtitle(doc, "（依据视频课题与《义务教育数学课程标准（2022 年版）》编写）")
    doc.add_paragraph()

    # 第一部分：视频分析报告
    add_h1(doc, "第一部分　视频分析报告")
    add_h2(doc, "一、视频基本信息")
    meta = [
        ("视频标题", v.get("title", "")),
        ("视频来源", "哔哩哔哩 (Bilibili)　" + v.get("url", "")),
        ("UP 主", v.get("uploader", "")),
        ("视频时长", v.get("duration_text", "")),
        ("画质规格", v.get("quality", "")),
        ("发布时间", v.get("publish", "")),
        ("互动数据", v.get("stats", "")),
        ("内容定位", "%s · %s · %s" % (m["subject"], m["lesson_type"], m["textbook"])),
    ]
    t = doc.add_table(rows=len(meta), cols=2)
    style_table(t)
    from docx.shared import Cm
    t.columns[0].width = Cm(3.5)
    t.columns[1].width = Cm(12.5)
    for i, (k, val) in enumerate(meta):
        set_cell(t.rows[i].cells[0], k, bold=True, color=NAVY)
        set_cell(t.rows[i].cells[1], str(val))

    add_h2(doc, "二、教学结构分析")
    add_para(doc, "说明：" + v.get("note", ""), color=GREY)
    for s in c.get("analysis", {}).get("structure", []):
        add_bullet(doc, "%s：%s" % (s["phase"], s["desc"]))

    add_h2(doc, "三、教学亮点与可借鉴点")
    for b in c.get("analysis", {}).get("highlights", []):
        add_bullet(doc, b)

    add_h2(doc, "四、使用说明")
    add_para(doc, "本教案依据视频课题、2022 版新课标与教材内容编写，可供教师备课、磨课、二次授课使用。"
                   "原视频未提供字幕，具体的课堂提问语、学生生成性回答与板书细节，建议结合视频实际观看后补充微调。")
    doc.add_page_break()

    # 第二部分：配套教案
    add_h1(doc, "第二部分　《%s》配套教案（详案）" % m["title"])
    add_h2(doc, "一、教学基本信息")
    info = [
        ("课题", m["title"] + "（比的意义）"),
        ("学段年级", m["grade"]),
        ("教材", m["textbook"]),
        ("课型", m["lesson_type"]),
        ("课时", m["duration"]),
        ("对应课标", m["standard"]),
    ]
    t = doc.add_table(rows=len(info), cols=2)
    style_table(t)
    from docx.shared import Cm
    t.columns[0].width = Cm(3.5)
    t.columns[1].width = Cm(12.5)
    for i, (k, val) in enumerate(info):
        set_cell(t.rows[i].cells[0], k, bold=True, color=NAVY)
        set_cell(t.rows[i].cells[1], str(val))

    add_h2(doc, "二、课标分析（2022 新课标）")
    add_para(doc, "“比”属于“数与代数”领域。新课标要求学生“经历从具体情境中抽象出比的过程，理解比的意义，"
                  "能解决按比分配的简单问题”。本课重点发展学生的符号意识、模型观念与应用意识，"
                  "并以“三会”为素养导向：会用数学的眼光观察、会用数学的思维思考、会用数学的语言表达现实世界。")
    add_h2(doc, "三、教材分析")
    add_para(doc, c.get("textbook_analysis", ""))
    add_h2(doc, "四、学情分析")
    add_para(doc, c.get("learning_analysis", ""))
    add_h2(doc, "五、教学目标")
    for i, o in enumerate(c.get("objectives", []), 1):
        add_bullet(doc, "%d. %s" % (i, o))
    add_h2(doc, "六、教学重难点")
    add_bullet(doc, "教学重点：" + c.get("keypoints", {}).get("focus", ""))
    for d in c.get("keypoints", {}).get("difficulties", []):
        add_bullet(doc, "教学难点：" + d)
    add_h2(doc, "七、教学准备")
    add_para(doc, c.get("prep", ""))

    add_h2(doc, "八、教学过程（详案）")

    def phase(title, minutes, rows):
        add_h2(doc, title + "（约 %d 分钟）" % minutes)
        tt = doc.add_table(rows=1, cols=3)
        style_table(tt)
        from docx.shared import Cm
        tt.columns[0].width = Cm(5.2)
        tt.columns[1].width = Cm(5.2)
        tt.columns[2].width = Cm(5.6)
        header_row(tt, ["教师活动", "学生活动", "设计意图"])
        for tea, stu, inten in rows:
            cells = tt.add_row().cells
            set_cell(cells[0], tea)
            set_cell(cells[1], stu)
            set_cell(cells[2], inten)

    for p in c.get("process", []):
        phase(p["name"], p.get("minutes", 0),
              [(p.get("teacher", ""), p.get("student", ""), p.get("intent", ""))])

    add_h2(doc, "九、板书设计")
    b = c.get("board", {})
    add_para(doc, b.get("meaning", ""), bold=True)
    add_para(doc, b.get("formula", ""))
    add_para(doc, b.get("parts", ""))
    add_para(doc, b.get("relation", ""))
    add_para(doc, "比　—　除法　—　分数（联系与区别）", bold=True)
    rel = b.get("rel_table", [])
    if rel:
        rt = doc.add_table(rows=len(rel), cols=4)
        style_table(rt)
        for ri, row in enumerate(rel):
            for ci, val in enumerate(row):
                if ri == 0:
                    set_cell(rt.rows[ri].cells[ci], val, bold=True,
                             color=__import__("docx").shared.RGBColor(0xFF, 0xFF, 0xFF))
                    shade(rt.rows[ri].cells[ci], "2E74B5")
                else:
                    set_cell(rt.rows[ri].cells[ci], val, bold=(ci == 0))

    add_h2(doc, "十、作业设计")
    for h in c.get("homework", []):
        add_bullet(doc, h)
    add_h2(doc, "十一、教学反思（课后填写）")
    add_para(doc, c.get("reflection", ""), color=GREY)
    doc.add_page_break()

    # 第三部分：AI 视频生成提示词
    add_h1(doc, "第三部分　配套 AI 视频生成提示词（讲解动画）")
    add_para(doc, "以下提示词用于生成约 30–45 秒的《%s》概念讲解动画，可作为课堂导入或课后微课素材。"
                  "采用“导演思维”编写，含主体、动作、场景、风格与运镜，并提供中英文对照，适配可灵/即梦/CogVideoX 等工具。"
                  % m["title"], color=GREY)
    for shot in c.get("ai_prompts", []):
        add_h2(doc, shot["title"])
        add_para(doc, "【中文】" + shot["cn"])
        add_para(doc, "【EN】" + shot["en"])
    add_para(doc, "生成建议：可逐镜生成后拼接为一段 30 秒左右的微课；建议统一“扁平教育插画”风格与 16:9 画幅以保证连贯；"
                  "人物若出现，建议用首帧锁定或参考图避免换脸。", color=GREY)

    doc.save(out_path)
    print("SAVED:", out_path)


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else DEFAULT
    c = load(src)
    out_dir = os.path.join(os.getcwd(), "output")
    os.makedirs(out_dir, exist_ok=True)
    out = os.path.join(out_dir, "%s_优质课配套教案.docx" % c["meta"]["title"])
    build(c, out)


if __name__ == "__main__":
    main()
