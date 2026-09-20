# -*- coding: utf-8 -*-
"""《视频转教案》Skill 共享的 DOCX 排版助手。

被 gen_lesson_plan.py / gen_script.py / gen_shuoke.py 调用，集中维护中文字体、
配色与表格样式，保证三套文档风格统一、可直接用 Word/WPS 二次编辑。

仅依赖 python-docx（运行环境：WorkBuddy 托管的 Python，已预装 python-docx）。
"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---- 统一配色 ----
NAVY = RGBColor(0x1F, 0x49, 0x7D)
BLUE = RGBColor(0x2E, 0x74, 0xB5)
GREY = RGBColor(0x60, 0x60, 0x60)
DARK = RGBColor(0x22, 0x22, 0x22)
GREEN = RGBColor(0x2E, 0x6B, 0x2E)
PURPLE = RGBColor(0x6A, 0x0D, 0xAD)
ORANGE = RGBColor(0xC0, 0x55, 0x00)

BODY_FONT = "宋体"
HEAD_FONT = "黑体"


def set_cjk(run, font=BODY_FONT, size=None, bold=None, color=None, italic=False):
    """设置中日韩字体并可选字号/加粗/颜色/斜体。"""
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), font)
    rfonts.set(qn('w:ascii'), font)
    rfonts.set(qn('w:hAnsi'), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if italic:
        run.font.italic = italic


def shade(cell, hex_fill):
    """给单元格填充背景色。"""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), hex_fill)
    tcPr.append(shd)


def set_cell(cell, text, bold=False, size=10.5, color=None, align=None, font=BODY_FONT):
    cell.text = ""
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    r = p.add_run(text)
    set_cjk(r, font, size=size, bold=bold, color=color)


def style_table(table):
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER


def header_row(table, headers, fill="2E74B5"):
    for c, h in zip(table.rows[0].cells, headers):
        set_cell(c, h, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(c, fill)


def add_title(doc, text, size=18):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_cjk(r, HEAD_FONT, size=size, bold=True, color=NAVY)
    return p


def add_subtitle(doc, text, size=10.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_cjk(r, BODY_FONT, size=size, color=GREY)
    return p


def add_h1(doc, text, color=NAVY, size=15):
    p = doc.add_heading(level=1)
    r = p.add_run(text)
    set_cjk(r, HEAD_FONT, size=size, bold=True, color=color)
    return p


def add_h2(doc, text, color=BLUE, size=13):
    p = doc.add_heading(level=2)
    r = p.add_run(text)
    set_cjk(r, HEAD_FONT, size=size, bold=True, color=color)
    return p


def add_para(doc, text, size=11, bold=False, color=None, indent=False, align=None, italic=False, bold_lead=None):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    if indent:
        p.paragraph_format.left_indent = Cm(0.75)
    if bold_lead:
        r0 = p.add_run(bold_lead)
        r0.bold = True
        set_cjk(r0)
    r = p.add_run(text)
    set_cjk(r, BODY_FONT, size=size, bold=bold, color=color, italic=italic)
    return p


def add_bullet(doc, text, size=11, bold=False, lead=None):
    p = doc.add_paragraph(style='List Bullet')
    if lead:
        r0 = p.add_run(lead)
        r0.bold = True
        set_cjk(r0)
    r = p.add_run(text)
    set_cjk(r, BODY_FONT, size=size, bold=bold)
    return p


def new_doc(body_font=BODY_FONT, body_size=11):
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = body_font
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), body_font)
    normal.font.size = Pt(body_size)
    return doc
