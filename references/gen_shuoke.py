# -*- coding: utf-8 -*-
"""生成《说课稿》DOCX（面向教研说课/面试陈述，八模块完整）。

对应 Skill 步骤：调用「教师教案与课堂活动」后产出说课稿。
用法：
    python gen_shuoke.py [lesson_content.json]
输出：./output/<课题>_说课稿.docx
"""
import sys
import os
import json

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from docx_common import (new_doc, add_para, add_bullet, set_cjk, set_cell,
                         style_table, NAVY, GREEN, DARK)

DEFAULT = os.path.join(HERE, "lesson_content.json")


def load(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def heading(doc, text, level=1):
    from docx.shared import Pt
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    if level == 1:
        p.alignment = __import__("docx").enum.text.WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text)
        r.bold = True; r.font.size = Pt(14); r.font.color.rgb = NAVY
        set_cjk(r, "黑体")
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(2)
        r2 = p2.add_run("─" * 18)
        r2.font.color.rgb = NAVY; r2.font.size = Pt(9)
    else:
        r = p.add_run(text)
        r.bold = True; r.font.size = Pt(12); r.font.color.rgb = GREEN
        set_cjk(r, "黑体")
    return p


def build(c, out_path):
    doc = new_doc()
    m = c["meta"]
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_before = Pt(4); t.paragraph_format.space_after = Pt(2)
    r = t.add_run("《%s》说课稿" % m["title"])
    r.bold = True; r.font.size = Pt(22); r.font.color.rgb = NAVY; set_cjk(r, "黑体")

    sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(2)
    r = sub.add_run("%s · %s · 对标 2022 年版义务教育数学课程标准" % (m["textbook"], m["grade"]))
    r.font.size = Pt(11); r.font.color.rgb = __import__("docx").shared.RGBColor(0x66, 0x66, 0x66); set_cjk(r)

    info = doc.add_paragraph(); info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.paragraph_format.space_after = Pt(8)
    r = info.add_run("说课人：          （请填写）    学科：%s    年级：%s" % (m["subject"], m["grade"]))
    r.font.size = Pt(10); set_cjk(r)

    sk = c.get("shuoke", {})
    cb = doc.add_paragraph(); cb.paragraph_format.space_after = Pt(6)
    r = cb.add_run("【一句话结论】"); r.bold = True; r.font.color.rgb = GREEN; set_cjk(r)
    r = cb.add_run(sk.get("one_liner", "")); set_cjk(r)

    # 一、说教材
    heading(doc, "一、说教材（教材分析）")
    add_para(doc, c.get("textbook_analysis", ""))
    add_para(doc, "本课属于“数与代数”领域下的“数量关系”主题，其核心是帮助学生建立“两个数相除又叫做两个数的比”"
                  "这一数学模型，体会比既可以表示同类量之间的倍数关系，也可以表示不同类量之间的相依关系（如路程∶时间 = 速度）。",
              bold_lead="")
    add_para(doc, "教材编排遵循“情境引入—意义建构—沟通联系—应用拓展”的线索，据此设计五个教学环节，把教材的“静态知识”转化为学生可参与的“动态探究”。",
              bold_lead="教材处理：")

    # 二、说学情
    heading(doc, "二、说学情（学情分析）")
    add_bullet(doc, c.get("learning_analysis", ""), lead="已有基础与特点：")
    add_bullet(doc, "对“比与除法、分数区别”“不同类量相比得到新量”“后项为何为0”需要结合意义才能想通。", lead="认知难点：")
    add_para(doc, "教学策略：用大情境降低抽象坡度，用对比表厘清易混点，用分层任务照顾差异——既不拔高要求，也不降低锚点目标。",
             bold_lead="应对策略：")

    # 三、说目标
    heading(doc, "三、说教学目标（核心素养导向）")
    add_para(doc, "基于 2022 年版课标“三会”核心素养，制定教学目标：")
    leads = ["1. 会用数学的眼光观察现实世界：", "2. 会用数学的思维思考现实世界：",
             "3. 会用数学的语言表达现实世界：", "4. 情感态度与价值观："]
    for i, o in enumerate(c.get("objectives", [])):
        add_bullet(doc, o, lead=leads[i] if i < len(leads) else "")

    # 四、说重难点
    heading(doc, "四、说教学重难点")
    add_bullet(doc, c.get("keypoints", {}).get("focus", ""), lead="教学重点：")
    for d in c.get("keypoints", {}).get("difficulties", []):
        add_bullet(doc, d, lead="教学难点：")

    # 五、说教法学法
    heading(doc, "五、说教法学法")
    heading(doc, "1. 教法", level=2)
    for item in sk.get("methods", {}).get("教法", []):
        add_bullet(doc, item["text"], lead=item["lead"])
    heading(doc, "2. 学法", level=2)
    for item in sk.get("methods", {}).get("学法", []):
        add_bullet(doc, item["text"], lead=item["lead"])

    # 六、说教学过程
    heading(doc, "六、说教学过程（约 45 分钟）")
    add_para(doc, "整节课按“引—探—练—结—延”五段推进，下面说清每一环节的设计意图。")
    for idx, p in enumerate(c.get("process", []), 1):
        label = ["一", "二", "三", "四", "五"][idx - 1] if idx <= 5 else str(idx)
        heading(doc, "（%s）%s（约 %d 分钟）" % (label, p["name"], p.get("minutes", 0)),
                level=2)
        # 用教师活动作为说课陈述
        add_para(doc, p.get("teacher", ""))
        add_para(doc, "设计意图：" + p.get("intent", ""), bold_lead="设计意图：")

    # 七、说板书设计
    heading(doc, "七、说板书设计")
    add_para(doc, "板书力求“脉络清晰、重点突出、助记助构”：")
    b = c.get("board", {})
    bt = doc.add_table(rows=6, cols=2)
    style_table(bt)
    from docx.shared import Cm
    bt.rows[0].cells[0].width = Cm(5.5); bt.rows[0].cells[1].width = Cm(11.0)
    data = [
        ("比的认识", "—— 两个数相除，又叫做两个数的比"),
        ("同类量：长∶宽", b.get("formula", "")),
        ("不同类量：路程∶时间", "90∶60 = 1.5（得到新量：速度）"),
        ("各部分名称", b.get("parts", "")),
        ("比、除法、分数", b.get("relation", "")),
        ("关键", "比的后项不能为 0（除数不能为 0）"),
    ]
    for i, (a, bb) in enumerate(data):
        c0 = bt.cell(i, 0); c0.text = ""
        rp = c0.paragraphs[0].add_run(a); rp.bold = True; rp.font.size = Pt(10.5); set_cjk(rp)
        c1 = bt.cell(i, 1); c1.text = ""
        rp2 = c1.paragraphs[0].add_run(bb); rp2.font.size = Pt(10.5); set_cjk(rp2)

    # 八、说特色与反思
    heading(doc, "八、说教学特色与反思")
    for f in sk.get("features", []):
        add_bullet(doc, f)
    add_para(doc, "可能的不足与改进：" + "；".join(sk.get("improve", [])), bold_lead="反思：")

    end = doc.add_paragraph(); end.paragraph_format.space_before = Pt(10)
    r = end.add_run("以上是我对《%s》一课的说课内容，恳请各位评委、老师批评指正。谢谢！" % m["title"])
    r.italic = True; r.font.size = Pt(11); set_cjk(r)

    doc.save(out_path)
    print("SAVED:", out_path)


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else DEFAULT
    c = load(src)
    out_dir = os.path.join(os.getcwd(), "output")
    os.makedirs(out_dir, exist_ok=True)
    out = os.path.join(out_dir, "%s_说课稿.docx" % c["meta"]["title"])
    build(c, out)


if __name__ == "__main__":
    main()
